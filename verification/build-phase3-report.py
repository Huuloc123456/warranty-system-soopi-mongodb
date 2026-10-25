"""Append executed Phase 3 evidence while retaining the original report."""
import json
from pathlib import Path
from datetime import datetime, timezone
from copy import deepcopy
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/phase3'
SOURCE = ROOT / 'artifacts/phase3/original.docx'
TARGET = OUT / 'Group2_QuanLyTrungTamBaoHanhDienTu_Phase3_KiemThuThucTe.docx'
result = json.loads((ROOT / 'verification/phase3-results.json').read_text(encoding='utf-8'))
snapshot = json.loads((OUT / 'mongodb-snapshot.json').read_text(encoding='utf-8'))
doc = Document(SOURCE)
original_paras = [p.text for p in doc.paragraphs]
original_tables = [[ [c.text for c in r.cells] for r in t.rows] for t in doc.tables]
next_page = False

def font(run, size=13, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        run._element.get_or_add_rPr().rFonts.set(qn('w:' + attr), 'Times New Roman')

def para(text='', size=13, bold=False, keep=False, center=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.keep_with_next = keep
    pf.widow_control = True
    pf.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    font(p.add_run(text), size, bold)
    return p

def heading(text, level=2):
    global next_page
    p = para(text, 14 if level == 1 else 13, True, True, level == 1)
    p.paragraph_format.page_break_before = next_page
    next_page = False
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(8)
    # Explicit outline levels make the appended chapter navigable in Word.
    ol = OxmlElement('w:outlineLvl'); ol.set(qn('w:val'), str(level - 1))
    p._p.get_or_add_pPr().append(ol)
    return p

def page():
    global next_page
    next_page = True

def table(headers, rows, widths, size=11):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    for i, s in enumerate(headers): t.rows[0].cells[i].text = str(s)
    for values in rows:
        cells = t.add_row().cells
        for i, s in enumerate(values): cells[i].text = str(s)
    for ri, row in enumerate(t.rows):
        rp = row._tr.get_or_add_trPr()
        rp.append(OxmlElement('w:cantSplit'))
        if ri == 0: rp.append(OxmlElement('w:tblHeader'))
        for ci, c in enumerate(row.cells):
            c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + edge)
                for a, v in [('val','single'), ('sz','4'), ('color','808080')]: e.set(qn('w:'+a),v)
                borders.append(e)
            pr.append(borders)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'E7E6E6'); pr.append(sh)
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = False
                for r in p.runs: font(r, size, ri == 0)
    spacer = para('', 1)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.space_after = Pt(3)
    return t

ui = [
    dict(id='UI-01', title='Danh sách khách hàng', steps='Đăng nhập admin, mở Quản trị → Tài khoản khách hàng.', expected='Có 3 hồ sơ mới, đúng họ tên/số điện thoại; chưa có tài khoản.', actual='KH-000002–000004 hiển thị đúng; tổng 4 hồ sơ gồm khách demo cũ.', status='PASS', image='01-khach-hang.png'),
    dict(id='UI-02', title='Danh sách và trạng thái phiếu', steps='Đăng nhập dispatcher, lọc số điện thoại 090902843; xem từng phiếu.', expected='Có 3 phiếu; laptop và iPhone đã bàn giao, Galaxy chưa phân công.', actual='Bộ lọc trả 3 phiếu; trạng thái và chi tiết khớp kết quả API.', status='PASS', image='02-ba-phieu.png'),
    dict(id='UI-03', title='Tra cứu công khai và chi phí', steps='Mở /portal; nhập TN-2026-1009-00012 và 0909028432; tra cứu, xem chi phí.', expected='Đã bàn giao; tổng 1.595.000đ gồm linh kiện, công và VAT.', actual='Đã bàn giao; 1.300.000 + 150.000 + 145.000 = 1.595.000đ.', status='PASS', image='03-thanh-toan-ban-giao.png'),
    dict(id='UI-04', title='Tồn kho sau xuất', steps='Đăng nhập warehouse; lọc SKU P3-MV0ODFHB.', expected='Pin 5, màn hình 4, đã giữ 0; đúng kệ P3-A1/P3-A2.', actual='Hai SKU hiển thị tồn 5 và 4; giữ 0; đúng kệ.', status='PASS', image='06-ton-kho.png'),
    dict(id='UI-05', title='Báo cáo SLA hàng đợi', steps='Đăng nhập dispatcher; mở Báo cáo → Phân bố SLA.', expected='Một phiếu đang xử lý, đúng hạn; không có phiếu trễ.', actual='Đúng hạn 1 (100%); sắp tới hạn 0; trễ hạn 0.', status='PASS', image='05-bao-cao-sla.png'),
    dict(id='UI-06', title='Tên nhóm Laptop', steps='Quan sát nhóm thiết bị của laptop và nhóm linh kiện pin Dell trong kho.', expected='Hiển thị “Laptop” tương ứng categoryCode=LAPTOP.', actual='Hiển thị “—”; dữ liệu API/MongoDB vẫn là LAPTOP. Ghi nhận F-01.', status='FAIL', image='06-ton-kho.png'),
]
(OUT / 'browser-checks.json').write_text(json.dumps({'campaign':result['campaign'], 'recordedAt':datetime.now(timezone.utc).isoformat(), 'cases':ui}, ensure_ascii=False, indent=2), encoding='utf-8')

page()
heading('CHƯƠNG 4. KIỂM THỬ HỆ THỐNG', 1)
para('Kiểm thử thực tế trên project warranty-system-soopi-mongodb', bold=True, center=True)
para('Ngày thực hiện: 09/10/2026. Mã đợt kiểm thử: P3-MV0ODFHB. Kịch bản API được ghi vào nhật ký trước khi chạy; kết quả, mã HTTP và thời gian được bổ sung ngay sau từng ca. Phần này được viết tiếp vào báo cáo Phase 3 hiện có.')
heading('4.1. Mục tiêu, môi trường và phương pháp')
para('Mục tiêu là kiểm chứng toàn bộ luồng tiếp nhận → phân công → chẩn đoán → sửa chữa → KCS (QC) → thanh toán → bàn giao, đồng thời kiểm tra phân quyền, dữ liệu trùng, IMEI, tồn kho và tra cứu khách hàng. Các thao tác nghiệp vụ được thực thi qua API của ứng dụng đang chạy; sau đó kiểm tra giao diện bằng trình duyệt và đọc dữ liệu đã lưu trong MongoDB.')
table(['Thành phần', 'Môi trường thực hiện'], [
    ('Ứng dụng', 'http://localhost:8081; project warranty-system-soopi-mongodb'),
    ('Backend', '.NET 10; MongoDB.Driver 3.12'),
    ('Cơ sở dữ liệu', 'MongoDB cục bộ; replica set rs0; database soopi_mongodb; 26 collections'),
    ('Thời gian chạy API', '09/10/2026, 14:59:59–15:03:06 (UTC+7)'),
    ('Đối chiếu MongoDB', '09/10/2026, 15:07:52 (UTC+7)'),
    ('Tài khoản kiểm thử', '8 tài khoản hiện có: admin, receptionist, dispatcher, technician, technician2, warehouse, cashier, 0900000001; mật khẩu demo chung 1234'),
    ('Công cụ', 'Script Node.js gọi HTTP thật; trình duyệt Codex; mongosh đọc đối chiếu'),
], [1.45,4.5])
para('Dữ liệu được tạo mới và lưu thật trong database để có thể mở lại, tiếp tục thao tác và kiểm tra. Tên, số điện thoại, địa chỉ, serial/IMEI và chữ ký là dữ liệu mẫu sát nghiệp vụ, không lấy từ khách hàng thật. Email dùng miền example.invalid; không gửi email/SMS hoặc thu tiền ngoài hệ thống.')
para('Điều kiện: ứng dụng và MongoDB đang chạy; các vai trò demo, danh mục hãng/loại thiết bị và nhà cung cấp NCC-001 đã tồn tại. Giữ nguyên dữ liệu cũ; chỉ bổ sung dữ liệu mang mã đợt P3-MV0ODFHB.')

page(); heading('4.2. Bộ dữ liệu nghiệp vụ đã tạo')
heading('4.2.1. Hồ sơ khách hàng',3)
table(['Mã khách', 'Họ tên / điện thoại', 'Địa chỉ mẫu'], [(c['code'], c['fullName']+'\n'+c['phone'], c['address']) for c in result['data']['customers']], [1.05,2.05,2.85])
para('Ba hồ sơ mới đang ACTIVE và chưa có tài khoản đăng nhập. Khách vẫn tra cứu bằng mã phiếu cùng số điện thoại. Các tài khoản demo cũ vẫn được giữ nguyên.')
heading('4.2.2. Thiết bị, hạn bảo hành và tình huống',3)
for i,t in enumerate(result['data']['tickets']):
    d=t['device']; c=t['customer']
    heading(f"{d['code']} — {d['product']['name']}",3)
    para(f"Khách {c['fullName']} ({c['code']}); mã phiếu {t['code']}. Serial/IMEI: {d['serialOrImei']}. Bảo hành: {d['warrantyActivatedOn']} đến {d['warrantyExpiresOn']}.", size=12)
    para([
        'Laptop không khởi động khi dùng pin, đèn nguồn chớp. Chẩn đoán lỏng đầu nối pin; vệ sinh, cố định đầu nối và cập nhật BIOS. Kiểm định lần đầu FAIL, khắc phục rồi kiểm định PASS. Bảo hành miễn phí; đã bàn giao.',
        'Màn hình sọc xanh, cảm ứng không ổn định. Thiết bị hết bảo hành; thay màn OLED. Khách đồng ý báo giá tại quầy; xuất 1 màn hình, kiểm định PASS, thanh toán tiền mặt và bàn giao.',
        'Sạc USB-C chập chờn. Thiết bị còn bảo hành; mới tiếp nhận, chưa phân công. Giữ phiếu trong hàng đợi để tiếp tục kiểm thử điều phối.',
    ][i], size=12)

page(); heading('4.2.3. Linh kiện và chứng từ phát sinh',3)
table(['SKU / linh kiện', 'Nhập / xuất / tồn', 'Giá / vị trí'], [
    ('P3-MV0ODFHB-1\nPin Dell Latitude 5420 63Wh', 'Nhập 5; xuất 0; tồn 5; giữ 0', 'Vốn 650.000đ\nDịch vụ 900.000đ\nKệ P3-A1'),
    ('P3-MV0ODFHB-2\nMàn hình Apple iPhone 13 OLED', 'Nhập 5; xuất 1; tồn 4; giữ 0', 'Vốn 850.000đ\nDịch vụ 1.300.000đ\nKệ P3-A2'),
], [2.45,1.7,1.8])
table(['Chứng từ', 'Mã / kết quả đã lưu'], [
    ('Phiếu nhập', result['data']['receipt']+'; APPROVED; 5 pin và 5 màn hình'),
    ('Báo giá iPhone', result['data']['quote']+'; APPROVED; khách ACCEPT'),
    ('Phiếu xuất', result['data']['issue']+'; 1 màn hình từ P3-A2'),
    ('Thanh toán laptop', 'PT-2026-00011; 0đ; FREE_WARRANTY / NONE'),
    ('Thanh toán iPhone', 'PT-2026-00012; 1.595.000đ; CHARGED / CASH'),
    ('Bàn giao', 'TN-2026-1009-00011 và TN-2026-1009-00012 đã DELIVERED; lưu xác nhận nghiệm thu và ảnh chữ ký mẫu'),
], [1.45,4.5])
para('Cách tính báo giá: màn hình 1.300.000đ + công 150.000đ = 1.450.000đ; VAT 10% = 145.000đ; tổng 1.595.000đ. Phần miễn phí không phát sinh xuất kho: pin Dell được nhập để kiểm tra tồn và dữ liệu danh mục.')
para('Đối chiếu trực tiếp database xác nhận hai phiếu DELIVERED, một phiếu RECEIVED; tồn pin 5, màn hình 4, reserved đều 0; hai khoản thanh toán lần lượt 0đ và 1.595.000đ. Không sử dụng số tồn lúc vừa tạo danh mục (onHand=0) làm số tồn cuối.')

for label, start, end in [
    ('4.3.1. Đăng nhập, danh mục và nhập kho',0,8),
    ('4.3.2. Luồng bảo hành miễn phí',8,18),
    ('4.3.3. Luồng sửa chữa có phí',18,28),
    ('4.3.4. Hàng đợi, báo cáo và kiểm tra dữ liệu',28,31),
]:
    page()
    if start == 0:
        heading('4.3. Bảng kịch bản và kết quả thực thi API')
        para('Mỗi dòng gồm thao tác/input, kết quả mong đợi trước khi chạy và kết quả quan sát thực tế. PASS của ca âm nghĩa là ứng dụng từ chối đúng yêu cầu; mã HTTP 400/401/403/409 được mong đợi trong các ca này. Các ca chạy tuần tự trên bộ dữ liệu ở mục 4.2; không phải toàn bộ quy trình được nhập bằng giao diện.', size=12)
    heading(label,3)
    rows=[]
    for c in result['cases'][start:end]:
        rows.append((c['id']+' — '+c['title']+'\n'+c['steps'], c['expected'], c['actual'], c['status']))
    table(['Ca / thao tác và dữ liệu', 'Mong đợi', 'Thực tế', 'KQ'], rows, [2.25,1.4,1.8,.5], 10.5)
    if start == 28:
        para('Báo cáo toàn database tại thời điểm chạy: tháng 10/2026 có 6 khoản thanh toán, tổng 4.895.000đ; trạng thái 12 DELIVERED và 1 RECEIVED. Các số này bao gồm dữ liệu demo cũ. Riêng đợt P3-MV0ODFHB phát sinh 1.595.000đ, hai phiếu bàn giao và một phiếu đang tiếp nhận.',size=12)

page(); heading('4.4. Kiểm tra giao diện trên trình duyệt')
para('Sau lượt API, đăng nhập các vai trò quản trị, điều phối và thủ kho trên cùng trang đăng nhập; mở các màn hình thực tế và tra cứu công khai. Kết quả dưới đây chỉ áp dụng cho các bước giao diện đã quan sát, không thay thế kiểm thử đầy đủ mọi chức năng UI.', size=12)
table(['Ca / thao tác', 'Mong đợi', 'Quan sát / kết quả'], [(c['id']+' — '+c['title']+'\n'+c['steps'],c['expected'],c['actual']+'\n'+c['status']) for c in ui],[2.2,1.5,2.25],10.5)

figures=[
    ('01-khach-hang.png','Hình 4.1. Hồ sơ khách hàng mới trong trang quản trị','UI-01 / TC-03: ba hồ sơ KH-000002, KH-000003, KH-000004 được hiển thị với tên và số điện thoại mẫu. Dòng khách demo cũ vẫn còn; ba hồ sơ mới chưa có tài khoản đăng nhập.'),
    ('02-ba-phieu.png','Hình 4.2. Bộ lọc trả về ba phiếu của đợt kiểm thử','UI-02 / TC-29: tìm tiền tố số điện thoại 090902843 trả 3 phiếu. Phần đang hiển thị trong ảnh là Galaxy S23, phiếu TN-2026-1009-00013, chưa phân công. Hai phiếu còn lại đã bàn giao theo nhật ký và dữ liệu MongoDB.'),
    ('04-tra-cuu-khach.png','Hình 4.3. Tra cứu công khai phiếu iPhone đã bàn giao','UI-03 / TC-28: nhập đúng mã TN-2026-1009-00012 và số điện thoại 0909028432. Kết quả hiển thị iPhone 13 và nhãn “Đã bàn giao”; không cần tạo thêm tài khoản khách hàng.'),
    ('03-thanh-toan-ban-giao.png','Hình 4.4. Chi phí trên trang tra cứu khách hàng','UI-03 / TC-23, TC-26: màn hình hiển thị linh kiện 1.300.000đ, công 150.000đ, VAT 145.000đ, tổng 1.595.000đ. Thanh toán CASH được xác nhận riêng bằng chứng từ PT-2026-00012 trong nhật ký API và MongoDB.'),
    ('06-ton-kho.png','Hình 4.5. Tồn kho sau khi xuất một màn hình','UI-04, UI-06 / TC-24: pin còn 5 tại P3-A1, màn hình còn 4 tại P3-A2; đã giữ đều 0. Ảnh cũng ghi nhận lỗi F-01: cột nhóm thiết bị của pin Dell hiển thị “—” thay vì “Laptop”.'),
    ('05-bao-cao-sla.png','Hình 4.6. Báo cáo SLA sau lượt kiểm thử','UI-05: còn một phiếu đang xử lý và đúng hạn (100%); sắp tới hạn 0, trễ hạn 0. Đây là ảnh màn hình báo cáo thực tế sau khi hai phiếu đã bàn giao và một phiếu được giữ trong hàng đợi.'),
]
for i,(file,title,caption) in enumerate(figures):
    page()
    if i == 0: heading('4.5. Ảnh minh chứng kiểm thử thực tế')
    heading(title,3)
    para(caption,size=12)
    img=OUT/'screenshots'/file
    with Image.open(img) as im: w,h=im.size
    width=min(5.85,6.65*w/h)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(6)
    p.add_run().add_picture(str(img),width=Inches(width))

page(); heading('4.6. Tổng hợp kết quả và lỗi ghi nhận')
table(['Nhóm kiểm thử', 'Đã chạy', 'Đạt', 'Không đạt'],[
    ('API / nghiệp vụ (kết quả cuối)',31,31,0),
    ('Giao diện quan sát trên trình duyệt',6,5,1),
    ('Tổng ca trong báo cáo',37,36,1),
],[2.75,1.1,1.05,1.05])
para('Nhật ký có 92 lượt HTTP, gồm lời gọi nghiệp vụ, ca âm, khôi phục phiên và các lần chạy chưa hoàn tất. Con số này không có nghĩa 92 lời gọi đều trả HTTP 200. Hai lần thử ban đầu được lưu riêng trong attempts; sau khi sửa dữ liệu/kịch bản, 31 ca API có kết quả cuối PASS.')
heading('4.6.1. Lỗi giao diện F-01',3)
table(['Thuộc tính', 'Nội dung'],[
    ('Hiện tượng', 'Nhóm Laptop hiển thị “—” ở danh sách phiếu và dòng linh kiện pin Dell trong kho.'),
    ('Tái hiện', 'Mở phiếu laptop TN-2026-1009-00011 hoặc /warehouse, lọc P3-MV0ODFHB; quan sát nhóm của SKU P3-MV0ODFHB-1.'),
    ('Mong đợi / thực tế', 'Mong đợi: Laptop. Thực tế: —. API và database vẫn có categoryCode=LAPTOP.'),
    ('Mức ảnh hưởng', 'Nhẹ: tên nhóm hiển thị thiếu; số lượng tồn và các nghiệp vụ đã kiểm thử vẫn đúng.'),
    ('Đề xuất / trạng thái', 'Kiểm tra ánh xạ mã LAPTOP sang nhãn tiếng Việt ở giao diện; chưa sửa trong đợt kiểm thử này.'),
],[1.45,4.5],11)
heading('4.6.2. Hai lần chỉnh dữ liệu và kịch bản kiểm thử',3)
para('TC-19 ban đầu dùng IMEI 15 chữ số nhưng sai checksum Luhn; ứng dụng trả 400 IMEI_INVALID đúng quy tắc. Đã sửa bộ sinh IMEI và tạo thành công 359915328831249. TC-23 ban đầu gọi nhầm route /tickets/{code}/quotation-decision và nhận 404; đã chuyển sang /quotations/{code}/customer-decision, dùng lại báo giá đã tạo và hoàn tất. Đây là lỗi fixture/kịch bản, không được tính là lỗi nghiệp vụ của ứng dụng.',size=12)

page(); heading('4.7. Phạm vi kết luận và hướng kiểm thử tiếp')
para('Trong phạm vi đã chạy, ứng dụng thực hiện đúng hai luồng sửa chữa, phân quyền, chặn dữ liệu trùng, xác thực IMEI, xuất kho, ghi thanh toán và bàn giao. Giao diện còn lỗi nhãn Laptop F-01. Kết quả chưa đủ để kết luận mọi chức năng của hệ thống đều đạt.')
para('Chưa kiểm thử tải đồng thời, mất kết nối/failover MongoDB, thanh toán ngân hàng/QR thật, gửi OTP/SMS/email thật và tất cả chức năng khác trong phần phân tích thiết kế. Thanh toán trong đợt này là chứng từ CASH mẫu trong hệ thống; chữ ký bàn giao là ảnh kiểm thử được gắn nhãn rõ ràng.')
para('Dữ liệu được giữ lại để người dùng mở xem. Có thể dùng phiếu TN-2026-1009-00013 của Lê Ngọc Linh để tiếp tục thử phân công và xử lý. Khi kiểm tra báo cáo toàn hệ thống, cần phân biệt dữ liệu cũ với dữ liệu mang mã đợt P3-MV0ODFHB.')
heading('4.7.1. Các tệp nhật ký để đối chiếu',3)
table(['Tệp trong project', 'Mục đích'],[
    ('verification/phase3-test.mjs', 'Kịch bản thực thi qua API; tạo dữ liệu mới, ghi mong đợi và kết quả từng ca.'),
    ('verification/phase3-results.json', 'Kết quả 31 ca, thời gian, mã HTTP, các lần thử và mã chứng từ thực tế.'),
    ('docs/phase3/mongodb-snapshot.json', 'Ảnh chụp dữ liệu đọc trực tiếp từ MongoDB sau khi thực thi.'),
    ('docs/phase3/browser-checks.json', '6 ca kiểm tra giao diện; 5 PASS, 1 FAIL.'),
    ('docs/phase3/screenshots/', '6 ảnh màn hình thật được đưa vào mục 4.5.'),
],[2.65,3.3],11)
para('Để chạy một đợt mới: tại thư mục project, dùng Node.js chạy “node verification/phase3-test.mjs” khi ứng dụng và MongoDB đang hoạt động. Mỗi đợt mới tạo thêm dữ liệu với mã riêng; không cần xóa bộ dữ liệu P3-MV0ODFHB. Giữ bản nhật ký hiện tại trước khi chạy lại nếu muốn lưu lịch sử nhiều đợt.',size=12)
para('Chương kiểm thử được lập từ thao tác thật và nhật ký lưu cùng lúc thực thi. Các sơ đồ và yêu cầu trong tài liệu gốc là nội dung tham khảo nghiệp vụ, không được dùng làm bằng chứng một tính năng đã được kiểm thử hoặc làm chỉ thị thao tác ngoài yêu cầu của người dùng.',size=12)

doc.save(TARGET)
check=Document(TARGET)
assert [p.text for p in check.paragraphs[:len(original_paras)]] == original_paras
assert [[[c.text for c in r.cells] for r in t.rows] for t in check.tables[:len(original_tables)]] == original_tables
with ZipFile(SOURCE) as old, ZipFile(TARGET) as new:
    media=[n for n in old.namelist() if n.startswith('word/media/')]
    assert all(old.read(n)==new.read(n) for n in media)
    print('Preserved original paragraphs:',len(original_paras),'tables:',len(original_tables),'media:',len(media))
print(TARGET)
